"""Turn the audit sample into a spreadsheet a person can actually sit and fill.

The markdown sheet was hard to read on a phone, which is where this is most
likely to get done. Same twenty rows, same seed, same bands -- just legible.

    python discovery-engine/audit/make_xlsx.py      # write audit_sample.xlsx
    python discovery-engine/audit/make_audit.py --score   # read it back

Column F is the only one that has to be filled: yes or no, from a dropdown.
G and H are there if you want to say what the right label was, which makes the
disagreements usable rather than just countable.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from make_audit import BANDS, FIELDS, load, sample

HERE = Path(__file__).resolve().parent
OUT = HERE / "audit_sample.xlsx"

HEAD = PatternFill("solid", fgColor="1A73E8")
BAND_FILL = {"A": PatternFill("solid", fgColor="E8F0FE"),
             "B": PatternFill("solid", fgColor="FEF7E0"),
             "C": PatternFill("solid", fgColor="FCE8E6")}
COLS = [("#", 5), ("Band", 7), ("What the post says", 78),
        ("AI: relevant?", 13), ("AI: labels", 30),
        ("Agree? (yes/no)", 15), ("If no, the right label is", 30), ("Notes", 26)]


def build():
    rows = sample(load())
    wb = Workbook()
    ws = wb.active
    ws.title = "Hand audit"

    ws["A1"] = ("Twenty AI classifications to check by hand. Read the post, then put yes or no "
                "in 'Agree?'. Bands B and C are rows the model REJECTED — they are here so the "
                "audit can catch what the filter threw away, not only what it wrongly kept.")
    ws["A1"].font = Font(size=11, italic=True, color="5F6368")
    ws.merge_cells("A1:H1")
    ws.row_dimensions[1].height = 42
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="center")

    for i, (name, width) in enumerate(COLS, start=1):
        c = ws.cell(row=2, column=i, value=name)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[2].height = 28
    ws.freeze_panes = "A3"

    for n, r in enumerate(rows, start=1):
        row = n + 2
        text = " ".join(r["full_text"].split()) or "(no text)"
        labels = "\n".join(f"{f}: {r[f]}" for f in FIELDS if f != "is_relevant")
        vals = [n, r["_band"], text, "yes" if r["is_relevant"] else "no", labels, "", "", ""]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(row=row, column=i, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if i <= 5:
                c.fill = BAND_FILL[r["_band"]]
        ws.cell(row=row, column=6).font = Font(bold=True)
        ws.row_dimensions[row].height = max(46, min(150, 13 * (len(text) // 70 + 2)))

    dv = DataValidation(type="list", formula1='"yes,no"', allow_blank=True)
    dv.prompt, dv.promptTitle = "Would you have filed it the same way?", "yes or no"
    ws.add_data_validation(dv)
    dv.add(f"F3:F{len(rows) + 2}")

    last = len(rows) + 2
    ws[f"A{last + 2}"] = "Agreed:"
    ws[f"B{last + 2}"] = f'=COUNTIF(F3:F{last},"yes")'
    ws[f"A{last + 3}"] = "Disagreed:"
    ws[f"B{last + 3}"] = f'=COUNTIF(F3:F{last},"no")'
    ws[f"A{last + 4}"] = "Still blank:"
    ws[f"B{last + 4}"] = f'={len(rows)}-COUNTIF(F3:F{last},"yes")-COUNTIF(F3:F{last},"no")'
    for k in range(2, 5):
        ws[f"A{last + k}"].font = Font(bold=True)

    key = [""] + [f"Band {b}: {n} rows — {label}" for b, n, label, _ in BANDS]
    for j, line in enumerate(key):
        ws.cell(row=last + 6 + j, column=1, value=line).font = Font(size=10, color="5F6368")

    wb.save(OUT)
    return OUT, len(rows)


if __name__ == "__main__":
    p, n = build()
    print(f"wrote {p.relative_to(HERE.parents[1])} — {n} rows")
