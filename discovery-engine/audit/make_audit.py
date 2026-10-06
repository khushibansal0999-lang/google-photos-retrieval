"""Draw a stratified sample of AI classifications for a human to check by hand.

The brief asks for this directly: "Manually audit 15-20 AI classifications to
demonstrate quality control." The point is not to prove the model was right.
It is to show the headline numbers survived a human looking at them.

The sample is deliberately NOT drawn only from the rows the model said yes to.
A sample of accepted rows can only ever find false positives, which flatters
the filter: it cannot tell you what the filter threw away. So the 20 rows span
three bands, and two of them are the model's rejections:

   10  problem_class = vague_memory_retrieval   the 84 the whole deck rests on
    5  relevant, but classed as something else  did we miss real cases?
    5  is_relevant = false                      did the filter bin a real one?

Seeded, so the same 20 rows come back every run and the audit is reproducible.

    python discovery-engine/audit/make_audit.py          # write the blank sheet
    python discovery-engine/audit/make_audit.py --score  # score it once filled
"""
import argparse
import csv
import json
import random
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
PROC = ROOT / "discovery-engine" / "data" / "processed"
HERE = Path(__file__).resolve().parent
SHEET = HERE / "audit_sample.md"
SEED = 20261003

BANDS = [
    ("A", 10, "Classed as a vague-memory retrieval failure (the headline 84)",
     lambda r: r["problem_class"] == "vague_memory_retrieval"),
    ("B", 5, "Judged relevant, but classed as a different kind of problem",
     lambda r: r["is_relevant"] and r["problem_class"] != "vague_memory_retrieval"),
    ("C", 5, "Judged NOT relevant, and dropped before analysis",
     lambda r: not r["is_relevant"]),
]

# The four labels worth checking. Anything more and nobody finishes the audit.
FIELDS = ["is_relevant", "problem_class", "failure_stage", "cues_remembered"]


def load():
    full = {}
    for line in (PROC / "tagged.jsonl").open():
        d = json.loads(line)
        full[d["id"]] = (d.get("text") or "").strip()
    rows = []
    for r in csv.DictReader((PROC / "tagged_flat.csv").open()):
        r["is_relevant"] = r["is_relevant"].strip().lower() in ("true", "1")
        r["full_text"] = full.get(r["id"], r["text"])
        rows.append(r)
    return rows


def sample(rows):
    rng = random.Random(SEED)
    out = []
    for band, n, label, keep in BANDS:
        pool = [r for r in rows if keep(r)]
        picked = rng.sample(pool, min(n, len(pool)))
        for r in picked:
            r["_band"], r["_band_label"] = band, label
            r["_pool"] = len(pool)
        out.extend(picked)
    return out


def write_sheet(picked):
    L = ["# Hand audit of 20 AI classifications", "",
         "The discovery engine tagged 1,196 posts with Gemini. These 20 were drawn",
         "at random, with a fixed seed, from three bands -- including two bands of",
         "rows the model **rejected**, so the audit can catch what the filter threw",
         "away and not only what it wrongly kept.", "",
         "For each row: read the post, then put `yes` or `no` in **Agree**. If no,",
         "say what the right label was. That is the whole task.", ""]
    for band, n, label, _ in BANDS:
        pool = next(r["_pool"] for r in picked if r["_band"] == band)
        L.append(f"- **Band {band}** ({n} of {pool}): {label}")
    L += ["", "---", ""]

    for i, r in enumerate(picked, 1):
        txt = " ".join(r["full_text"].split())
        if len(txt) > 700:
            txt = txt[:700] + " …"
        L += [f"## {i}. Band {r['_band']} · `{r['id']}` · {r['source']}", "",
              "> " + (txt or "*(no text)*"), "",
              "| field | what the AI said |", "|---|---|"]
        for f in FIELDS:
            L.append(f"| {f} | `{r[f]}` |")
        L += ["",
              "**Agree:**  ",
              "**If no, the right label is:**  ", "", "---", ""]

    L += ["## Result", "",
          "Fill these in once every row above has an answer, then put the",
          "agreement line on slide 1.", "",
          "- Rows audited: 20",
          "- Rows where the human agreed with every field: __ of 20",
          "- Fields corrected: __",
          "- What the disagreements had in common: ____", ""]
    SHEET.write_text("\n".join(L))
    return SHEET


XLSX = HERE / "audit_sample.xlsx"


def _score_xlsx():
    """Read column F of the spreadsheet. Returns None if it is not there."""
    if not XLSX.exists():
        return None
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("(found audit_sample.xlsx but openpyxl is not installed; "
              "pip install openpyxl)")
        return None
    ws = load_workbook(XLSX, data_only=True).active
    yes = no = blank = 0
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row, min_col=1, max_col=6):
        if not isinstance(row[0].value, int):
            break                       # past the twenty rows, into the totals
        v = str(row[5].value or "").strip().lower()
        yes += v.startswith("y")
        no += v.startswith("n")
        blank += not v
    return yes, no, blank


def score():
    """Count the yes/no answers, from the spreadsheet if there is one."""
    counts = _score_xlsx()
    if counts is None:
        if not SHEET.exists():
            raise SystemExit("Nothing to score yet. Run without --score first.")
        yes = no = blank = 0
        for line in SHEET.read_text().splitlines():
            if line.startswith("**Agree:**"):
                v = line.split("**Agree:**", 1)[1].strip().strip("*").lower()
                if v.startswith("y"):
                    yes += 1
                elif v.startswith("n"):
                    no += 1
                else:
                    blank += 1
    else:
        yes, no, blank = counts
        print(f"(read {XLSX.name})")
    done = yes + no
    print(f"agreed {yes}, disagreed {no}, not yet filled {blank}")
    if done:
        print(f"agreement: {yes}/{done} = {100*yes/done:.0f}%")
        print(f'\nfor slide 1: "{yes} of {done} hand-checked classifications agreed."')
    return yes, no, blank


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--score", action="store_true")
    a = ap.parse_args()
    if a.score:
        score()
    else:
        rows = load()
        picked = sample(rows)
        p = write_sheet(picked)
        print(f"wrote {p.relative_to(ROOT)} -- {len(picked)} rows")
        for band, n, _, _ in BANDS:
            got = sum(1 for r in picked if r["_band"] == band)
            print(f"  band {band}: {got}")
